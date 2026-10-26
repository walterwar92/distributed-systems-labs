# -*- coding: utf-8 -*-
"""Собирает быстрый отчёт в формате .docx.

Запуск:  python build_docx.py   ->  ОТЧЁТ.docx
"""
import os
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "ОТЧЁТ.docx")

BODY = "Cambria"
MONO = "Consolas"

INK = RGBColor(0x1C, 0x1C, 0x1A)
ACCENT = RGBColor(0x23, 0x40, 0x8E)   # глубокий индиго для заголовков
SOFT = RGBColor(0x5A, 0x5A, 0x54)

FILL_CODE = "F4F4F0"
FILL_HEAD = "E8E8E2"

WIDGET_URL = "https://claude.ai/code/artifact/939246f8-6910-41a9-8687-2feded11c211"


# --------------------------------------------------------------------------
# низкоуровневые помощники: python-docx не умеет заливку и рамки напрямую
# --------------------------------------------------------------------------
def shade(element, fill):
    """Заливка ячейки таблицы или абзаца."""
    pr = element.get_or_add_tcPr() if element.tag.endswith("tc") else element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pr.append(shd)


def style_run(run, name=BODY, size=10.5, bold=False, italic=False, color=INK):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    # Кириллица берёт шрифт из w:cs/w:hAnsi — задаём явно, иначе Word подставит свой.
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        fonts.set(qn(attr), name)


def para(doc, text="", size=10.5, bold=False, italic=False, color=INK,
         space_before=0, space_after=6, align=None, indent=None):
    p = doc.add_paragraph()
    fmt = p.paragraph_format
    fmt.space_before = Pt(space_before)
    fmt.space_after = Pt(space_after)
    fmt.line_spacing = 1.15
    if align is not None:
        p.alignment = align
    if indent is not None:
        fmt.left_indent = Cm(indent)
    if text:
        style_run(p.add_run(text), size=size, bold=bold, italic=italic, color=color)
    return p


def rich(doc, chunks, size=10.5, space_after=6, indent=None):
    """Абзац из кусочков (текст, жирный?) — для выделений внутри строки."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if indent is not None:
        p.paragraph_format.left_indent = Cm(indent)
    for text, bold in chunks:
        style_run(p.add_run(text), size=size, bold=bold)
    return p


def heading(doc, text, level=1):
    sizes = {0: 16, 1: 13, 2: 11.5}
    p = doc.add_paragraph()
    fmt = p.paragraph_format
    fmt.space_before = Pt(2 if level == 0 else (16 if level == 1 else 12))
    fmt.space_after = Pt(6 if level == 0 else 6)
    fmt.keep_with_next = True
    style_run(
        p.add_run(text),
        size=sizes[level],
        bold=True,
        color=INK if level == 0 else ACCENT,
    )
    if level <= 1:
        borders = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "4")
        bottom.set(qn("w:color"), "C9C9C2")
        borders.append(bottom)
        p._element.get_or_add_pPr().append(borders)
    return p


def bullets(doc, items, size=10.5):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.12
        if isinstance(item, (list, tuple)):
            for text, bold in item:
                style_run(p.add_run(text), size=size, bold=bold)
        else:
            style_run(p.add_run(item), size=size)


def code_block(doc, lines, size=9):
    """Моноширинный блок в одной залитой ячейке — рамка не рвётся между строк."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    cell = table.rows[0].cells[0]
    shade(cell._element, FILL_CODE)
    cell.width = Cm(17)

    cell.paragraphs[0]._element.getparent().remove(cell.paragraphs[0]._element)
    for line in lines:
        p = cell.add_paragraph()
        fmt = p.paragraph_format
        fmt.space_before = Pt(0)
        fmt.space_after = Pt(0)
        fmt.line_spacing = 1.0
        style_run(p.add_run(line if line else " "), name=MONO, size=size)
    para(doc, space_after=4)
    return table


def table_block(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT

    head = table.rows[0]
    for cell, text, width in zip(head.cells, headers, widths):
        shade(cell._element, FILL_HEAD)
        cell.width = Cm(width)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        style_run(p.add_run(text), size=9.5, bold=True)

    for row_data in rows:
        row = table.add_row()
        for cell, text, width in zip(row.cells, row_data, widths):
            cell.width = Cm(width)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.08
            style_run(p.add_run(text), size=9.5)
    para(doc, space_after=6)
    return table


def figure(doc, path, caption, width_cm=15.5):
    if not os.path.exists(path):
        para(doc, f"[нет файла {os.path.basename(path)} — запустите plot_regions.py]",
             italic=True, color=SOFT)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(path, width=Cm(width_cm))
    cap = para(doc, caption, size=9, italic=True, color=SOFT,
               align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
    return cap


# --------------------------------------------------------------------------
def build():
    doc = Document()

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    normal = doc.styles["Normal"]
    normal.font.name = BODY
    normal.font.size = Pt(10.5)

    # ---------------- шапка ----------------
    heading(doc, "Классификация ДУЧП 2-го порядка и приведение "
                 "к каноническому виду", level=0)
    rich(doc, [("Состав группы: ", True), ("______________________________", False)],
         size=10.5, space_after=2)
    para(doc,
         "Выполнены все три пункта — A, B и C. Ядро написано на Python 3, "
         "виджет — на JavaScript. Готовые функции, решающие саму задачу, "
         "не использовались; что именно сделано вручную — в конце отчёта.",
         color=SOFT, space_after=4)

    # ---------------- A ----------------
    heading(doc, "A. Определение типа и замены переменных по коэффициентам")
    para(doc, "Реализованы два независимых метода — это заодно даёт взаимную проверку.")

    heading(doc, "A.1. Произвольное число переменных — метод Лагранжа", level=2)
    para(doc,
         "Уравнение Σ aᵢⱼ·u_xᵢxⱼ + F = 0 задаёт характеристическую квадратичную форму "
         "Q(λ) = Σ aᵢⱼ λᵢλⱼ. Выделением полных квадратов она приводится к виду Σ βᵢ μᵢ² "
         "заменой μ = Cλ. Далее P = C⁻¹, а замена переменных в самом уравнении есть "
         "ξ = Pᵀx — уравнение принимает вид Σ βᵢ·u_ξᵢξᵢ + F̃ = 0.")
    para(doc, "Тип читается по набору β:", space_after=4)
    table_block(
        doc,
        ["Условие на β", "Тип уравнения"],
        [
            ["все βᵢ ≠ 0, знаки не все одинаковы", "гиперболический"],
            ["есть βᵢ = 0", "параболический"],
            ["все βᵢ ≠ 0 и одного знака", "эллиптический"],
        ],
        [9.5, 7.0],
    )
    para(doc,
         "Тонкое место — форма без единого квадрата (например Q = λ₁λ₂): выделять нечего. "
         "Тогда сначала делается замена λᵢ = (νᵢ+νⱼ)/2, λⱼ = (νᵢ−νⱼ)/2, и только потом "
         "выделяются квадраты. Это ровно приём из Примера 1 разбора к лекции.")

    rich(doc, [("Пример 1. ", True), ("u_xy − 2u_xz + u_yz + u_x + ½u_y = 0", False)],
         space_after=3)
    code_block(doc, [
        "тип:  гиперболический,   beta = (1, -1, 1)",
        "",
        "      [ 1/2   1/2  -1/2 ]           [ 1   1  -V2/2 ]",
        "C  =  [ 1/2  -1/2   3/2 ]     P  =  [ 1  -1   V2   ]",
        "      [  0     0    V2  ]           [ 0   0   V2/2 ]",
        "",
        "замена:  xi1 = x + y     xi2 = x - y     xi3 = V2*(-x + 2y + z)/2",
        "канон:   u_xi1xi1 - u_xi2xi2 + u_xi3xi3 + F = 0",
        "проверка P^T*A*P = diag(beta):  сходится        (V2 = корень из 2)",
    ])

    rich(doc, [("Пример 2. ", True), ("u_xx + 2u_xy + 5u_yy − 32u = 0", False)],
         space_after=3)
    code_block(doc, [
        "тип:  эллиптический,   beta = (1, 1)",
        "замена:  xi1 = x        xi2 = -x/2 + y/2",
        "канон:   u_xi1xi1 + u_xi2xi2 + F = 0",
    ])
    para(doc,
         "Оба ответа совпали с разбором к лекции посимвольно, включая радикалы.",
         italic=True, color=SOFT)

    heading(doc, "A.2. Две переменные — метод характеристик", level=2)
    para(doc,
         "Для a₁₁u_xx + 2a₁₂u_xy + a₂₂u_yy + F = 0 характеристическое уравнение "
         "a₁₁dy² − 2a₁₂dxdy + a₂₂dx² = 0 имеет дискриминант D = a₁₂² − a₁₁a₂₂:")
    bullets(doc, [
        "D > 0 — две характеристики, ξ = φ, η = ψ, канонический вид u_ξη = Φ;",
        "D = 0 — одна характеристика, ξ = φ, η — любая с ненулевым якобианом, u_ηη = Φ;",
        "D < 0 — корни комплексно сопряжены, ξ = Re φ, η = Im φ, u_ξξ + u_ηη = Φ.",
    ])
    para(doc, space_after=4)
    table_block(
        doc,
        ["Уравнение", "D", "Тип", "Замена", "Новые A₁₁, A₁₂, A₂₂"],
        [
            ["u_xx − 2u_xy + u_yy + 9u_x + 9u_y − 9u = 0", "0",
             "параболический", "ξ = x+y, η = x", "0, 0, 1"],
            ["волновое  u_tt − u_xx = 0", "1",
             "гиперболический", "ξ = y−x, η = y+x", "0, −2, 0"],
            ["теплопроводности  u_t = u_xx", "0",
             "параболический", "ξ = y, η = x", "0, 0, 1"],
            ["Лапласа  u_xx + u_yy = 0", "−1",
             "эллиптический", "ξ = y, η = x", "1, 0, 1"],
        ],
        [5.0, 1.0, 3.1, 3.3, 3.6],
    )
    para(doc,
         "Отдельно разобран случай a₁₁ = 0, на котором формула dy/dx = (a₁₂ ± √D)/a₁₁ "
         "делится на ноль: характеристическое уравнение решается относительно dx/dy, "
         "а при a₁₁ = a₂₂ = 0 уравнение уже в первой канонической форме.")

    rich(doc, [("Как программа себя проверяет. ", True),
               ("Найденные ξ и η подставляются обратно, и пересчитываются "
                "коэффициенты при старших производных:", False)], space_after=3)
    code_block(doc, [
        "A11 = a11*xi_x^2 + 2*a12*xi_x*xi_y + a22*xi_y^2",
        "A12 = a11*xi_x*eta_x + a12*(xi_x*eta_y + xi_y*eta_x) + a22*xi_y*eta_y",
        "A22 = a11*eta_x^2 + 2*a12*eta_x*eta_y + a22*eta_y^2",
    ])
    para(doc,
         "Для гиперболического типа обязаны занулиться A₁₁ и A₂₂, для параболического — "
         "A₁₁ и A₁₂, для эллиптического — A₁₂ при A₁₁ = A₂₂. Это не пересказ теории, "
         "а вычисляемый признак: именно он стоит в тестах.")

    # ---------------- B ----------------
    heading(doc, "B. Интерактивный виджет")
    para(doc,
         "Файл widget/index.html — открывается двойным щелчком, интернет и установка "
         "не нужны. Опубликованная версия:", space_after=2)
    para(doc, WIDGET_URL, size=9.5, color=ACCENT, space_after=8)

    para(doc, "Три режима:", space_after=4)
    bullets(doc, [
        [("Две переменные. ", True),
         ("Ползунки и числовые поля для a₁₁, a₁₂, a₂₂. При любом движении сразу "
          "пересчитываются D, тип, характеристики, замена ξ и η, якобиан, новые "
          "коэффициенты и канонический вид. Показан и ход решения.", False)],
        [("Переменные коэффициенты. ", True),
         ("Коэффициенты вводятся как выражения от x, y (y, x*y, 1-x^2, sqrt, sin). "
          "Строится карта областей; при наведении курсора виден тип в точке, "
          "по щелчку — полный разбор в ней.", False)],
        [("n переменных. ", True),
         ("Матрица aᵢⱼ до 5×5: метод Лагранжа, β, тип, замена ξ = Pᵀx и результат "
          "проверки PᵀAP = diag(β).", False)],
    ])
    para(doc, space_after=2)
    para(doc,
         "В каждом режиме есть готовые примеры: волновое, теплопроводности, Лапласа, "
         "Трикоми, все примеры из лекции. Виджет не обращается к Python — вся "
         "арифметика продублирована на JavaScript, причём в режиме n переменных она "
         "точная: дроби на BigInt, без потери разрядов. Разбор выражений a(x, y) — "
         "собственный рекурсивный парсер, без eval.")

    # ---------------- C ----------------
    heading(doc, "C. Уравнения с переменными коэффициентами: где какой тип")
    para(doc,
         "Если a₁₁, a₁₂, a₂₂ зависят от точки, то и D(x, y) = a₁₂² − a₁₁a₂₂ зависит "
         "от точки. Плоскость разбивается на области постоянства знака D, и в каждой "
         "уравнение приводится к своему каноническому виду.")
    para(doc,
         "Алгоритм: вычисляем D символьно → находим границы D = 0 → берём в каждой "
         "области пробную точку → по знаку D определяем тип → приводим к каноническому "
         "виду с учётом знака.")
    para(doc,
         "Знак переменной приходится сообщать явно: при y < 0 подстановка y = −p с "
         "положительным p меняет знак дифференциала (dy = −dp), поэтому в "
         "характеристическом уравнении у a₁₂ меняется знак. Без этого радикалы вида "
         "√(−y) не упрощаются.")

    heading(doc, "Уравнение Трикоми  y·u_xx + u_yy = 0", level=2)
    para(doc, "D = −y, и уравнение меняет тип трижды:", space_after=4)
    table_block(
        doc,
        ["Область", "Тип", "Замена", "Новые A₁₁, A₁₂, A₂₂", "Канонический вид"],
        [
            ["y < 0", "гиперболический",
             "ξ = x + ⅔(−y)³ᐟ², η = x − ⅔(−y)³ᐟ²", "0, 2y, 0", "u_ξη = Φ"],
            ["y = 0", "параболический", "ξ = x, η = y", "0, 0, 1", "u_ηη = Φ"],
            ["y > 0", "эллиптический", "ξ = x, η = ⅔·y³ᐟ²", "y, 0, y",
             "u_ξξ + u_ηη = Φ"],
        ],
        [1.8, 3.2, 5.0, 3.0, 3.0],
    )
    para(doc, "Совпадает со слайдами 15–17 лекции.", italic=True, color=SOFT)
    figure(doc, os.path.join(HERE, "regions_tricomi.png"),
           "Области уравнения Трикоми; тонкие линии — характеристики "
           "x ± ⅔(−y)³ᐟ² = C в гиперболической области.")

    heading(doc, "Другие примеры", level=2)
    table_block(
        doc,
        ["Уравнение", "D", "Области"],
        [
            ["u_xx + x·u_yy = 0", "−x",
             "x < 0 гипербол., x = 0 парабол., x > 0 эллипт."],
            ["u_xx + xy·u_yy = 0", "−xy", "тип меняется по квадрантам"],
            ["(1 − x²)·u_xx + u_yy = 0", "x² − 1",
             "x < −1 гипербол., x = −1 парабол., −1 < x < 1 эллипт., "
             "x = 1 парабол., x > 1 гипербол."],
        ],
        [4.6, 1.8, 10.2],
    )
    para(doc,
         "Последний пример показывает, что областей может быть больше трёх: границ "
         "столько, сколько корней у D.")
    figure(doc, os.path.join(HERE, "regions_gallery.png"),
           "Карты областей: знак дискриминанта D = a₁₂² − a₁₁a₂₂ задаёт тип в каждой точке.",
           width_cm=16.0)

    # ---------------- проверка ----------------
    heading(doc, "Проверка корректности")
    rich(doc, [("python -m unittest discover -s tests -t .  →  ", False),
               ("96 тестов, все проходят.", True)])
    para(doc,
         "Тесты писались до кода, а эталонами служили ответы из лекции и разбора к ней: "
         "матрицы C и P Примера 1 сверяются посимвольно вместе с √2, характеристики "
         "Трикоми — с точностью до знака (φ и −φ задают одно семейство).")
    para(doc,
         "Помимо сверки с лекцией проверяются вычисляемые тождества, которые "
         "не зависят от способа решения:", space_after=4)
    bullets(doc, [
        "PᵀAP = diag(β) — для каждого разобранного уравнения;",
        "зануление нужных A₁₁, A₁₂, A₂₂ после подстановки найденной замены;",
        "ненулевой якобиан замены — иначе замена вырождена;",
        "согласие двух методов: для девяти наборов коэффициентов метод Лагранжа "
        "и метод характеристик обязаны дать один и тот же тип.",
    ])
    para(doc, space_after=2)
    para(doc,
         "Отдельно покрыты случаи, на которых наивная реализация ломается: a₁₁ = 0, "
         "форма без квадратов, пустая старшая часть, нулевая форма.")

    # ---------------- про встроенные функции ----------------
    heading(doc, "Про запрет встроенных функций")
    para(doc,
         "Готовые классификаторы ДУЧП и решатели уравнений не используются. "
         "Руками написаны:", space_after=4)
    bullets(doc, [
        "метод Лагранжа, включая замену при отсутствии квадратов;",
        "обращение матрицы методом Гаусса-Жордана над точными числами "
        "(не numpy.linalg);",
        "правило классификации по сигнатуре;",
        "решение характеристического уравнения, включая вырождение a₁₁ = 0;",
        "интегрирование характеристик разделением переменных;",
        "разбиение плоскости на области и обход связных компонент для подписей;",
        "в виджете — точная арифметика дробей и парсер выражений.",
    ])
    para(doc, space_after=2)
    para(doc,
         "sympy используется только как калькулятор точной арифметики — дроби, "
         "радикалы, производные, первообразные, корни многочлена. "
         "matplotlib — только рисование.")

    # ---------------- запуск ----------------
    heading(doc, "Как запустить")
    code_block(doc, [
        "python demo.py                               # все примеры из лекции",
        "python plot_regions.py                       # карты областей",
        "python -m unittest discover -s tests -t .    # тесты",
    ])
    para(doc, "Виджет — открыть widget/index.html в браузере.", space_after=0)

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(f"собрано: {path}  ({os.path.getsize(path)} байт)")
